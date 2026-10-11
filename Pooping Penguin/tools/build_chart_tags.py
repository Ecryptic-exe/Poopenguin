"""
Build data/chart_tags.json from the community chart spreadsheet (.xlsx).

    python tools/build_chart_tags.py                      (uses the .xlsx in tools/)
    python tools/build_chart_tags.py ChunithmChart.xlsx   (a name is looked up in tools/)
    python tools/build_chart_tags.py C:/full/path/to/sheet.xlsx

Needs openpyxl >= 3.1 (this tool only; the bot itself does not import it).
Re-run whenever the sheet is updated.

Reads sheets 13.7以下 / 13.7-13.9 / 14 / 14+ / 15以上 (not 刪除曲).
  A  Const   "■ 15.2"; the colour of the ■ is the difficulty
             (red = Expert, purple = Master, cyan = Ultima)
  B  Title
  C  Genre
  E  主要配置 -> tags, split on / (unweighted)
  F  SSS難度  -> number of stars (1-5)
  G  Chain    -> note count
  H  説明     -> note text (kept), plus flags from data/note_flags.json
             (regex per flag: tolerance, not_for_push, personal, ...)
Rows whose ■ has no usable colour default to Master and are listed at the end;
fix them in data/chart_overrides.json as {"Title|15.3": "Ultima"}.
"""
import json
import os
import re
import sys

TOOLS_DIR = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(TOOLS_DIR)
DEFAULT_XLSX = "ChunithmChart.xlsx"   # put the spreadsheet in vto2/tools/
sys.path.insert(0, ROOT)
from ratingbot.catalog import normalize   # noqa: E402

SHEETS = ["13.7以下", "13.7-13.9", "14", "14+", "15以上"]
DIFF_RGB = {"Expert": (255, 0, 0), "Master": (153, 0, 255), "Ultima": (0, 255, 255)}
MAX_COLOR_DIST = 160   # beyond this the colour is treated as "unknown"


def _nearest_difficulty(rgb):
    """'FFEA4335' -> 'Expert' (closest of the three marker colours) or None."""
    if not isinstance(rgb, str) or len(rgb) != 8:
        return None
    try:
        r, g, b = int(rgb[2:4], 16), int(rgb[4:6], 16), int(rgb[6:8], 16)
    except ValueError:
        return None
    best, dist = None, 1e9
    for name, (R, G, B) in DIFF_RGB.items():
        d = ((r - R) ** 2 + (g - G) ** 2 + (b - B) ** 2) ** 0.5
        if d < dist:
            best, dist = name, d
    return best if dist <= MAX_COLOR_DIST else None


def _find_xlsx(arg=None):
    """Resolve the spreadsheet: explicit path, a name inside tools/, or the
    only .xlsx sitting in tools/."""
    if arg:
        for cand in (arg, os.path.join(TOOLS_DIR, arg)):
            if os.path.isfile(cand):
                return cand
        sys.exit("Spreadsheet not found: %s\n(also looked in %s)" % (arg, TOOLS_DIR))
    default = os.path.join(TOOLS_DIR, DEFAULT_XLSX)
    if os.path.isfile(default):
        return default
    found = [f for f in os.listdir(TOOLS_DIR)
             if f.lower().endswith(".xlsx") and not f.startswith("~$")]
    if len(found) == 1:
        return os.path.join(TOOLS_DIR, found[0])
    sys.exit("Put the spreadsheet in %s as %s (or pass its path).%s" % (
        TOOLS_DIR, DEFAULT_XLSX,
        "" if not found else "\nFound several .xlsx files: " + ", ".join(found)))


def _tag_key(t):
    return re.sub(r"\s+", "", t).casefold()


def _load_flag_rules():
    path = os.path.join(ROOT, "data", "note_flags.json")
    rules = []
    if os.path.exists(path):
        with open(path, encoding="utf-8") as f:
            for flag, pat in json.load(f).items():
                if not flag.startswith("_"):
                    rules.append((flag, re.compile(pat)))
    return rules


def _load_aliases():
    path = os.path.join(ROOT, "data", "tag_aliases.json")
    lookup = {}
    if os.path.exists(path):
        with open(path, encoding="utf-8") as f:
            for canon, variants in json.load(f).items():
                if canon.startswith("_"):
                    continue
                lookup[_tag_key(canon)] = canon
                for v in variants:
                    lookup[_tag_key(v)] = canon
    return lookup


def main(xlsx, out_path):
    from openpyxl import load_workbook
    aliases = _load_aliases()
    flag_rules = _load_flag_rules()
    overrides = {}
    op = os.path.join(ROOT, "data", "chart_overrides.json")
    if os.path.exists(op):
        with open(op, encoding="utf-8") as f:
            overrides = json.load(f)

    wb = load_workbook(xlsx, rich_text=True)
    charts, defaulted, seen = [], [], set()
    for name in SHEETS:
        if name not in wb.sheetnames:
            print("skip missing sheet:", name)
            continue
        ws = wb[name]
        for r in range(2, ws.max_row + 1):
            v = ws.cell(r, 1).value
            title = ws.cell(r, 2).value
            if v is None or not title:
                continue
            rgb, text = None, str(v)
            if hasattr(v, "__iter__") and not isinstance(v, str):
                parts = list(v)
                text = "".join(getattr(p, "text", str(p)) for p in parts)
                for p in parts:
                    if "■" in getattr(p, "text", ""):
                        col = p.font.color if p.font else None
                        rgb = col.rgb if col is not None and isinstance(col.rgb, str) else None
            m = re.search(r"(\d+(?:\.\d+)?)", text)
            if not m:
                continue
            const = float(m.group(1))
            title = str(title).strip()
            diff = overrides.get("%s|%s" % (title, m.group(1))) or _nearest_difficulty(rgb)
            if diff is None:
                diff = "Master"
                defaulted.append("%s|%s  (%s row %d)" % (title, m.group(1), name, r))
            key = (normalize(title), diff)
            if key in seen:
                print("duplicate ignored:", title, diff)
                continue
            seen.add(key)

            tags = []
            raw = ws.cell(r, 5).value
            if raw:
                for t in re.split(r"[/／、,，]", str(raw)):
                    t = t.strip()
                    if t:
                        c = aliases.get(_tag_key(t), t)
                        if c not in tags:
                            tags.append(c)
            stars = ws.cell(r, 6).value
            stars = str(stars).count("★") if stars else 0
            notes = ws.cell(r, 7).value
            genre = ws.cell(r, 3).value
            note = ws.cell(r, 8).value
            note = " ".join(str(note).split()) if note else None
            flags = [f for f, rx in flag_rules if note and rx.search(note)]
            charts.append({
                "title": title,
                "key": normalize(title),
                "difficulty": diff,
                "const": const,
                "tags": tags,
                "sss_stars": stars or None,
                "notes": int(notes) if isinstance(notes, (int, float)) else None,
                "genre": str(genre).strip() if genre else None,
                "note": note,
                "flags": flags,
            })

    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump({"source": os.path.basename(xlsx), "charts": charts}, f,
                  ensure_ascii=False, indent=1)
    ge = [c for c in charts if c["const"] >= 14]
    from collections import Counter
    fc = Counter(f for c in charts for f in c["flags"])
    print("notes: %d charts, flags: %s" % (sum(1 for c in charts if c["note"]), dict(fc)))
    print("wrote %s: %d charts, %d at const>=14, %d of those tagged" % (
        out_path, len(charts), len(ge), sum(1 for c in ge if c["tags"])))
    if defaulted:
        print("\nNo usable difficulty colour -> defaulted to Master (override if wrong):")
        for d in defaulted:
            print("  ", d)


if __name__ == "__main__":
    main(_find_xlsx(sys.argv[1] if len(sys.argv) > 1 else None),
         sys.argv[2] if len(sys.argv) > 2
         else os.path.join(ROOT, "data", "chart_tags.json"))